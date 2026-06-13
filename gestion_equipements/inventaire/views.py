from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.forms import modelformset_factory
from django.http import HttpResponse

from .models import Inventaire, LigneInventaire
from .forms import InventaireForm, LigneInventaireForm
from salles.models import Salle
from equipements.models import Equipement

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
from datetime import date

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False


# ── Liste des inventaires ──────────────────────────────────────────────────────

@login_required
def liste_inventaires(request):
    inventaires = Inventaire.objects.select_related('cree_par').all()
    return render(request, 'inventaire/liste.html', {'inventaires': inventaires})


# ── Créer un inventaire ────────────────────────────────────────────────────────

@login_required
def creer_inventaire(request):
    if request.method == 'POST':
        form = InventaireForm(request.POST)
        if form.is_valid():
            inv = form.save(commit=False)
            inv.cree_par = request.user
            inv.save()
            messages.success(request, "Inventaire créé. Commencez la vérification salle par salle.")
            return redirect('inventaire:detail_inventaire', pk=inv.pk)
    else:
        form = InventaireForm()
    return render(request, 'inventaire/form.html', {'form': form, 'titre': 'Nouvel inventaire'})


# ── Détail inventaire (résumé par salle) ──────────────────────────────────────

@login_required
def detail_inventaire(request, pk):
    inventaire = get_object_or_404(Inventaire, pk=pk)
    salles = Salle.objects.filter(est_active=True)

    # Pour chaque salle : est-elle déjà vérifiée dans cet inventaire ?
    salles_data = []
    for salle in salles:
        lignes = inventaire.lignes.filter(salle=salle)
        salles_data.append({
            'salle': salle,
            'nb_verifies': lignes.count(),
            'nb_absents': lignes.filter(statut='absent').count(),
            'nb_endommages': lignes.filter(statut='endommage').count(),
            'faite': lignes.exists(),
        })

    return render(request, 'inventaire/detail.html', {
        'inventaire': inventaire,
        'salles_data': salles_data,
    })


# ── Saisie inventaire par salle ───────────────────────────────────────────────
@login_required
def saisie_salle(request, inv_pk, salle_pk):
    inventaire = get_object_or_404(Inventaire, pk=inv_pk)
    salle = get_object_or_404(Salle, pk=salle_pk)

    if inventaire.statut == 'cloture':
        messages.error(request, "Cet inventaire est clôturé.")
        return redirect('inventaire:detail_inventaire', pk=inv_pk)

    # Équipements empruntés activement par des étudiants de cette salle
    from emprunt.models import Mouvement
    emprunts_actifs = Mouvement.objects.filter(
        type_operation='emprunt',
        etat__in=['EN_COURS', 'VALIDE'],
        beneficiaire__type_contact='etudiant',
        beneficiaire__salle=salle
    ).select_related('equipement', 'beneficiaire')

    equipements = [m.equipement for m in emprunts_actifs]

    # Créer les lignes manquantes
    for eq in equipements:
        LigneInventaire.objects.get_or_create(
            inventaire=inventaire,
            salle=salle,
            equipement=eq,
            defaults={
                'statut': 'present',
                'quantite_attendue': 1,
                'quantite_constatee': 1,
            }
        )

    lignes = inventaire.lignes.filter(salle=salle).select_related('equipement')

    LigneFormSet = modelformset_factory(
        LigneInventaire,
        fields=['statut', 'quantite_constatee', 'observations'],
        extra=0,
        widgets={
            'observations': LigneInventaireForm.Meta.widgets['observations'],
        }
    )

    if request.method == 'POST':
        formset = LigneFormSet(request.POST, queryset=lignes)
        if formset.is_valid():
            instances = formset.save(commit=False)
            for inst in instances:
                inst.verifie_par = request.user
                inst.save()
            messages.success(request, f"Salle « {salle.nom} » vérifiée.")
            return redirect('inventaire:detail_inventaire', pk=inv_pk)
    else:
        formset = LigneFormSet(queryset=lignes)

    return render(request, 'inventaire/saisie_salle.html', {
        'inventaire': inventaire,
        'salle': salle,
        'formset': formset,
        'lignes': lignes,
        'paires': list(zip(lignes, formset)),
        'emprunts_actifs': emprunts_actifs,
    })


def saisie_salle(request, inv_pk, salle_pk):
    inventaire = get_object_or_404(Inventaire, pk=inv_pk)
    salle = get_object_or_404(Salle, pk=salle_pk)

    if inventaire.statut == 'cloture':
        messages.error(request, "Cet inventaire est clôturé.")
        return redirect('inventaire:detail_inventaire', pk=inv_pk)

    # Équipements empruntés activement par des étudiants de cette salle
    from emprunt.models import Mouvement
    emprunts_actifs = Mouvement.objects.filter(
        type_operation='emprunt',
        etat__in=['EN_COURS', 'VALIDE'],
        beneficiaire__type_contact='etudiant',
        beneficiaire__salle=salle
    ).select_related('equipement', 'beneficiaire')

    equipements = [m.equipement for m in emprunts_actifs]

    # Créer les lignes manquantes
    for eq in equipements:
        LigneInventaire.objects.get_or_create(
            inventaire=inventaire,
            salle=salle,
            equipement=eq,
            defaults={
                'statut': 'present',
                'quantite_attendue': 1,
                'quantite_constatee': 1,
            }
        )

    lignes = inventaire.lignes.filter(salle=salle).select_related('equipement')

    LigneFormSet = modelformset_factory(
        LigneInventaire,
        fields=['statut', 'quantite_constatee', 'observations'],
        extra=0,
        widgets={
            'observations': LigneInventaireForm.Meta.widgets['observations'],
        }
    )

    if request.method == 'POST':
        formset = LigneFormSet(request.POST, queryset=lignes)
        if formset.is_valid():
            instances = formset.save(commit=False)
            for inst in instances:
                inst.verifie_par = request.user
                inst.save()
            messages.success(request, f"Salle « {salle.nom} » vérifiée.")
            return redirect('inventaire:detail_inventaire', pk=inv_pk)
    else:
        formset = LigneFormSet(queryset=lignes)

    return render(request, 'inventaire/saisie_salle.html', {
        'inventaire': inventaire,
        'salle': salle,
        'formset': formset,
        'lignes': lignes,
        'paires': list(zip(lignes, formset)),
        'emprunts_actifs': emprunts_actifs,
    })

    inventaire = get_object_or_404(Inventaire, pk=inv_pk)
    salle = get_object_or_404(Salle, pk=salle_pk)

    if inventaire.statut == 'cloture':
        messages.error(request, "Cet inventaire est clôturé.")
        return redirect('inventaire:detail_inventaire', pk=inv_pk)

    # Équipements assignés à cette salle
    equipements = Equipement.objects.filter(salle=salle, est_actif=True)

    # Créer les lignes manquantes (pré-remplissage)
    for eq in equipements:
        LigneInventaire.objects.get_or_create(
            inventaire=inventaire,
            salle=salle,
            equipement=eq,
            defaults={
                'statut': 'present',
                'quantite_attendue': 1,
                'quantite_constatee': 1,
            }
        )

    lignes = inventaire.lignes.filter(salle=salle).select_related('equipement')

    LigneFormSet = modelformset_factory(
        LigneInventaire,
        fields=['statut', 'quantite_constatee', 'observations'],
        extra=0,
        widgets={
            'observations': LigneInventaireForm.Meta.widgets['observations'],
        }
    )

    if request.method == 'POST':
        formset = LigneFormSet(request.POST, queryset=lignes)
        if formset.is_valid():
            instances = formset.save(commit=False)
            for inst in instances:
                inst.verifie_par = request.user
                inst.save()
            messages.success(request, f"Salle « {salle.nom} » vérifiée avec succès.")
            return redirect('inventaire:detail_inventaire', pk=inv_pk)
    else:
        formset = LigneFormSet(queryset=lignes)

    return render(request, 'inventaire/saisie_salle.html', {
        'inventaire': inventaire,
        'salle': salle,
        'formset': formset,
        'lignes': lignes,
        'paires': list(zip(lignes, formset)),
    })


# ── Clôturer un inventaire ────────────────────────────────────────────────────

@login_required
def cloturer_inventaire(request, pk):
    inventaire = get_object_or_404(Inventaire, pk=pk)
    if request.method == 'POST':
        inventaire.statut = 'cloture'
        inventaire.date_fin = date.today()
        inventaire.save()
        messages.success(request, "Inventaire clôturé.")
    return redirect('inventaire:detail_inventaire', pk=pk)


# ── Supprimer inventaire ───────────────────────────────────────────────────────

@login_required
def supprimer_inventaire(request, pk):
    inventaire = get_object_or_404(Inventaire, pk=pk)
    if request.method == 'POST':
        inventaire.delete()
        messages.success(request, "Inventaire supprimé.")
        return redirect('inventaire:liste_inventaires')
    return render(request, 'inventaire/confirmer_suppression.html', {'inventaire': inventaire})


# ══════════════════════════════════════════════════════════════════════════════
# RAPPORTS
# ══════════════════════════════════════════════════════════════════════════════

def _border():
    s = Side(style='thin', color='000000')
    return Border(left=s, right=s, top=s, bottom=s)


# ── Rapport inventaire Excel (toutes salles ou une salle) ─────────────────────

@login_required
def rapport_inventaire_excel(request, pk, salle_pk=None):
    inventaire = get_object_or_404(Inventaire, pk=pk)
    lignes = inventaire.lignes.select_related('equipement', 'salle', 'verifie_par')
    titre_fichier = f"inventaire_{inventaire.date_debut}"

    if salle_pk:
        salle = get_object_or_404(Salle, pk=salle_pk)
        lignes = lignes.filter(salle=salle)
        titre_fichier += f"_{salle.nom}"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventaire"

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1a3c5e")
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Titre
    ws.merge_cells('A1:H1')
    ws['A1'] = f"RAPPORT D'INVENTAIRE — {inventaire}"
    ws['A1'].font = Font(bold=True, size=13)
    ws['A1'].alignment = center

    ws.merge_cells('A2:H2')
    ws['A2'] = f"Période : {inventaire.date_debut} → {inventaire.date_fin or 'En cours'}"
    ws['A2'].alignment = center

    # En-têtes
    headers = ['Salle', 'Équipement', 'Référence', 'Statut',
               'Qté attendue', 'Qté constatée', 'Écart', 'Observations']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = _border()

    # Données
    COULEURS = {'present': 'C6EFCE', 'absent': 'FFC7CE', 'endommage': 'FFEB9C'}
    for row_idx, ligne in enumerate(lignes, 5):
        vals = [
            ligne.salle.nom,
            ligne.equipement.nom,
            ligne.equipement.reference or '—',
            ligne.get_statut_display(),
            ligne.quantite_attendue,
            ligne.quantite_constatee,
            ligne.ecart,
            ligne.observations or '',
        ]
        fill = PatternFill("solid", fgColor=COULEURS.get(ligne.statut, 'FFFFFF'))
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = _border()
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if col == 4:
                cell.fill = fill

    # Largeurs colonnes
    for col, w in zip('ABCDEFGH', [20, 25, 15, 12, 12, 12, 8, 30]):
        ws.column_dimensions[col].width = w

    # Stats résumé
    ws_stat = wb.create_sheet("Résumé")
    ws_stat['A1'] = "Résumé de l'inventaire"
    ws_stat['A1'].font = Font(bold=True, size=12)
    stats = [
        ('Total équipements vérifiés', lignes.count()),
        ('Présents', lignes.filter(statut='present').count()),
        ('Absents', lignes.filter(statut='absent').count()),
        ('Endommagés', lignes.filter(statut='endommage').count()),
    ]
    for i, (label, val) in enumerate(stats, 3):
        ws_stat[f'A{i}'] = label
        ws_stat[f'B{i}'] = val
        ws_stat[f'A{i}'].font = Font(bold=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{titre_fichier}.xlsx"'
    return response


# ── Rapport inventaire PDF (toutes salles ou une salle) ───────────────────────

@login_required
def rapport_inventaire_pdf(request, pk, salle_pk=None):
    if not REPORTLAB_OK:
        messages.error(request, "ReportLab n'est pas installé. Installez-le avec : pip install reportlab")
        return redirect('inventaire:detail_inventaire', pk=pk)

    inventaire = get_object_or_404(Inventaire, pk=pk)
    lignes = inventaire.lignes.select_related('equipement', 'salle').order_by('salle__nom', 'equipement__nom')
    titre_fichier = f"inventaire_{inventaire.date_debut}"

    if salle_pk:
        salle = get_object_or_404(Salle, pk=salle_pk)
        lignes = lignes.filter(salle=salle)
        titre_fichier += f"_{salle.nom}"

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    elements = []

    # Titre
    elements.append(Paragraph(f"<b>RAPPORT D'INVENTAIRE</b>", styles['Title']))
    elements.append(Paragraph(f"{inventaire}", styles['Normal']))
    elements.append(Paragraph(
        f"Période : {inventaire.date_debut} → {inventaire.date_fin or 'En cours'}",
        styles['Normal']
    ))
    elements.append(Spacer(1, 0.5*cm))

    # Tableau
    data = [['Salle', 'Équipement', 'Référence', 'Statut', 'Attendu', 'Constaté', 'Écart']]
    COULEURS_PDF = {
        'present':   colors.HexColor('#C6EFCE'),
        'absent':    colors.HexColor('#FFC7CE'),
        'endommage': colors.HexColor('#FFEB9C'),
    }
    row_colors = [('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a3c5e'))]

    for i, ligne in enumerate(lignes, 1):
        data.append([
            ligne.salle.nom,
            ligne.equipement.nom,
            ligne.equipement.reference or '—',
            ligne.get_statut_display(),
            str(ligne.quantite_attendue),
            str(ligne.quantite_constatee),
            str(ligne.ecart),
        ])
        c = COULEURS_PDF.get(ligne.statut, colors.white)
        row_colors.append(('BACKGROUND', (3, i), (3, i), c))

    table = Table(data, colWidths=[3*cm, 4.5*cm, 3*cm, 2.5*cm, 2*cm, 2.5*cm, 1.5*cm])
    style = TableStyle([
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
    ] + row_colors)
    table.setStyle(style)
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))

    # Résumé
    elements.append(Paragraph("<b>Résumé :</b>", styles['Normal']))
    elements.append(Paragraph(
        f"Total : {lignes.count()} | "
        f"Présents : {lignes.filter(statut='present').count()} | "
        f"Absents : {lignes.filter(statut='absent').count()} | "
        f"Endommagés : {lignes.filter(statut='endommage').count()}",
        styles['Normal']
    ))

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{titre_fichier}.pdf"'
    return response


# ── Rapport de perte Excel ────────────────────────────────────────────────────

@login_required
def rapport_perte_excel(request, pk):
    inventaire = get_object_or_404(Inventaire, pk=pk)
    lignes_absentes = inventaire.lignes.filter(
        statut='absent'
    ).select_related('equipement', 'salle')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport de perte"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="8B0000")
    center = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('A1:F1')
    ws['A1'] = f"RAPPORT DE PERTE — {inventaire}"
    ws['A1'].font = Font(bold=True, size=13)
    ws['A1'].alignment = center

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Date : {date.today().strftime('%d/%m/%Y')}"
    ws['A2'].alignment = center

    headers = ['Salle', 'Équipement', 'Référence', 'Catégorie', 'Qté perdue', 'Observations']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = _border()

    fill_rouge = PatternFill("solid", fgColor="FFC7CE")
    for row_idx, ligne in enumerate(lignes_absentes, 5):
        ecart_abs = abs(ligne.ecart) if ligne.ecart < 0 else ligne.quantite_attendue
        vals = [
            ligne.salle.nom,
            ligne.equipement.nom,
            ligne.equipement.reference or '—',
            ligne.equipement.categorie.nom if ligne.equipement.categorie else '—',
            ecart_abs,
            ligne.observations or '',
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = _border()
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.fill = fill_rouge

    for col, w in zip('ABCDEF', [20, 25, 15, 15, 12, 30]):
        ws.column_dimensions[col].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_perte_{inventaire.date_debut}.xlsx"'
    return response


# ── Rapport de perte PDF ──────────────────────────────────────────────────────

@login_required
def rapport_perte_pdf(request, pk):
    if not REPORTLAB_OK:
        messages.error(request, "ReportLab n'est pas installé.")
        return redirect('inventaire:detail_inventaire', pk=pk)

    inventaire = get_object_or_404(Inventaire, pk=pk)
    lignes_absentes = inventaire.lignes.filter(
        statut='absent'
    ).select_related('equipement', 'salle')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("<b>RAPPORT DE PERTE</b>", styles['Title']))
    elements.append(Paragraph(f"{inventaire}", styles['Normal']))
    elements.append(Paragraph(f"Généré le : {date.today().strftime('%d/%m/%Y')}", styles['Normal']))
    elements.append(Spacer(1, 0.5*cm))

    if not lignes_absentes.exists():
        elements.append(Paragraph("Aucune perte constatée lors de cet inventaire.", styles['Normal']))
    else:
        data = [['Salle', 'Équipement', 'Référence', 'Catégorie', 'Qté perdue', 'Observations']]
        for ligne in lignes_absentes:
            ecart_abs = abs(ligne.ecart) if ligne.ecart < 0 else ligne.quantite_attendue
            data.append([
                ligne.salle.nom,
                ligne.equipement.nom,
                ligne.equipement.reference or '—',
                ligne.equipement.categorie.nom if ligne.equipement.categorie else '—',
                str(ecart_abs),
                ligne.observations or '',
            ])

        table = Table(data, colWidths=[2.5*cm, 4*cm, 2.5*cm, 3*cm, 2*cm, 4.5*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#8B0000')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#FFC7CE')),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 0.5*cm))
        elements.append(Paragraph(
            f"<b>Total des pertes : {lignes_absentes.count()} équipement(s)</b>",
            styles['Normal']
        ))

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="rapport_perte_{inventaire.date_debut}.pdf"'
    return response